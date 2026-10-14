from dataclasses import dataclass, field
from pathlib import Path
import re
import openpyxl
from .normalizacion import norm, text, bank, identifier, match_id, BANK_CODES


def workbook(source):
    if hasattr(source, 'seek'):
        source.seek(0)
    return openpyxl.load_workbook(source, read_only=True, data_only=True)


def header_rows(sheet, required):
    rows = sheet.iter_rows(values_only=True)
    for _ in range(20):
        row = next(rows, ())
        names = [norm(v) for v in row]
        if required <= set(names):
            if len([n for n in names if n]) != len(set(n for n in names if n)):
                raise ValueError('Encabezados duplicados en ' + sheet.title)
            return {name: i for i, name in enumerate(names) if name}, rows
    raise ValueError('Faltan columnas: ' + ', '.join(sorted(required)))


def load_dictionary(source):
    w = workbook(source)
    mapping, invalid = {}, []
    try:
        headers, rows = header_rows(w.active, {'RNC34', 'CORRECTO'})
        entries = []
        for row in rows:
            original, correct = row[headers['RNC34']], row[headers['CORRECTO']]
            if not text(original):
                continue
            if not text(correct) or text(correct).startswith('#'):
                invalid.append(text(original))
                continue
            entries.append((norm(original), norm(correct)))
        for original, correct in entries + [(v, v) for _, v in entries]:
            if original in mapping and mapping[original] != correct:
                raise ValueError(f'Diccionario ambiguo: {original} tiene varios equivalentes.')
            mapping[original] = correct
        if not mapping:
            raise ValueError('El diccionario no contiene equivalencias válidas.')
        return mapping, invalid
    finally:
        w.close()


def load_billing(source, mapping):
    """Read the actual indented one-column agent -> banks hierarchy."""
    w = workbook(source)
    pairs, agent, duplicates = set(), None, 0
    try:
        for row in w.active.iter_rows():
            cells = [c for c in row if text(c.value)]
            if not cells:
                continue
            if len(cells) != 1:
                raise ValueError('AGREGADOS EN COBRO debe tener una columna jerárquica agente / bancos.')
            cell = cells[0]
            value = norm(cell.value)
            is_bank = bank(value) in BANK_CODES.values() or value.startswith('BANCO ')
            if value in mapping and not is_bank:
                if cell.alignment.indent:
                    raise ValueError(f'Agente con sangría de banco: {value}. Revise la jerarquía.')
                agent = mapping[value]
            elif is_bank:
                if agent is None:
                    raise ValueError(f'Banco sin agente anterior: {value}.')
                pair = (agent, bank(value))
                duplicates += pair in pairs
                pairs.add(pair)
            else:
                raise ValueError(f'Fila no reconocida en AGREGADOS EN COBRO: {value}. Actualice el diccionario.')
        if not pairs:
            raise ValueError('AGREGADOS EN COBRO no contiene combinaciones válidas.')
        return pairs, duplicates
    finally:
        w.close()


@dataclass
class Park:
    keys: set = field(default_factory=set)
    afipos: set = field(default_factory=set)
    fallback: dict = field(default_factory=dict)
    duplicates: int = 0
    usable: bool = True
    reason: str = ''


def load_park(source, mapping, usable=True, reason=''):
    if source is None:
        return Park(usable=False, reason='No se cargó PARQUE PARCIAL.')
    w = workbook(source)
    result = Park(usable=usable, reason=reason)
    try:
        if '7A_PARQUE_PARCIAL_ACTIVOS' not in w.sheetnames:
            raise ValueError('PARQUE PARCIAL no contiene la hoja 7A_PARQUE_PARCIAL_ACTIVOS.')
        h, rows = header_rows(w['7A_PARQUE_PARCIAL_ACTIVOS'], {'AFIPOS', 'AGREGADO', 'BANCO', 'TERMINAL'})
        for row in rows:
            if not any(text(v) for v in row):
                continue
            a = norm(row[h['AGREGADO']])
            if a not in mapping:
                raise ValueError(f'Agente del parque sin equivalencia: {a}.')
            afipos, terminal = match_id(row[h['AFIPOS']]), match_id(row[h['TERMINAL']])
            b = bank(row[h['BANCO']])
            if not afipos or not b:
                raise ValueError('El parque contiene claves de equipo incompletas.')
            key = (mapping[a], b, afipos, terminal)
            result.duplicates += key in result.keys
            result.keys.add(key)
            result.afipos.add(afipos)
            affiliate = match_id(row[h['AFILIADO']]) if 'AFILIADO' in h else ''
            if affiliate and terminal:
                result.fallback.setdefault((mapping[a], b, affiliate, terminal), set()).add(afipos)
        return result
    finally:
        w.close()


MONTHS = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']


def park_period_issue(filename, period):
    """A filename can disprove the period, never prove the file's contents."""
    year, month = map(int, period.split('-'))
    name = norm(Path(filename).stem).lower()
    found = [i for i, word in enumerate(MONTHS, 1) if word in name]
    if 'setiembre' in name:
        found.append(9)
    years = [int(y) for y in re.findall(r'\b20\d{2}\b', name)]
    numeric = re.findall(r'(20\d{2})[-_ ](0?[1-9]|1[0-2])(?:\D|$)', name)
    if (found and set(found) != {month}) or (years and set(years) != {year}) or any((int(y), int(m)) != (year, month) for y, m in numeric):
        return f'El nombre del parque ({filename}) no corresponde al período {period}. Casos especiales pendientes.'
    return ''
