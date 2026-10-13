import io
import tempfile
import unittest
from pathlib import Path
import pandas as pd
import openpyxl
from agregados.normalizacion import COLUMNS, MODES, bank, identifier
from agregados.entradas import Park, load_dictionary, load_billing, load_park, park_period_issue
from agregados.facturacion import classify
from agregados.procesamiento import process


def excel(rows, title='Hoja1'):
    f = io.BytesIO()
    w = openpyxl.Workbook()
    w.active.title = title
    for row in rows:
        w.active.append(row)
    w.save(f)
    f.seek(0)
    return f


def record(agent='VEPAGOS', bank='BANCO DE VENEZUELA', rango='TRANSANDO', tx='1', afipos='001234561', **extra):
    row = {c: '' for c in COLUMNS if c not in ['FACTURACION', 'PERTENENCIA2']}
    row.update(AFIPOS=afipos, TERMINAL='29600100123020303913 [00982] DSPREAD', SERIAL='29600100123020303913', PERTENENCIA=agent, NOMBRE_BANCO=bank, RANGO=rango,
               CANT_TRANS_ACUM_MES=tx, ANO_PROCESO='2026', MES_PROCESO='10', NUMPOS='1.0')
    row.update(extra)
    return row


MAPPING = {s: s for s in ['VEPAGOS', 'RAPIDPAGO', 'SERVICIOS J&G', 'DISGLOBAL', 'UBIIPAGOS', 'POSGLOBAL', 'LATIN PAGOS 2021', 'CARROPAGO', 'POSITIVO PLUS', 'ACIERTO TECNOLOGICO', 'AGILPAGO', 'PEQUENO']}
MAPPING.update({'RAPID PAGO': 'RAPIDPAGO', 'FLEXIPOS': 'POSITIVO PLUS', 'AGILPAGOS': 'AGILPAGO'})


class BillingTests(unittest.TestCase):
    def classify(self, rows, park=None, pairs=None):
        return classify(pd.DataFrame(rows), MAPPING, pairs or {('VEPAGOS', 'BANCO DE VENEZUELA')}, park or Park())[0]

    def test_agent_bank_pair_not_bank_alone(self):
        out = self.classify([record(), record('PEQUENO'), record(bank='MI BANCO')])
        self.assertEqual(out.FACTURACION.tolist(), [MODES[0], MODES[1], MODES[1]])

    def test_four_exempt_agents_only_when_zero_and_range_agrees(self):
        agents = ['DISGLOBAL', 'UBIIPAGOS', 'POSGLOBAL', 'LATIN PAGOS 2021']
        for agent in agents:
            out = self.classify([record(agent, rango='30 DIAS SIN TX', tx='0'), record(agent), record(agent, rango='TRANSANDO', tx='0')])
            self.assertEqual(out.FACTURACION.tolist(), [MODES[2], MODES[1], MODES[3]])
        self.assertEqual(self.classify([record(rango='30 DIAS SIN TX', tx='0')]).FACTURACION.iloc[0], MODES[0])

    def test_pos_indicator_conflict_is_reported(self):
        out = self.classify([record('DISGLOBAL', rango='120 DIAS SIN TX', tx='0', POS_CON_TRANSACCION='1')])
        self.assertEqual(out.FACTURACION.iloc[0], MODES[2])
        self.assertIn('POS_CON_TRANSACCION', out._ISSUE.iloc[0])

    def test_special_keys_require_agent_bank_afipos_numpos(self):
        for agent, b in [('RAPIDPAGO', 'BANCO DE VENEZUELA'), ('RAPIDPAGO', 'BANFANB'), ('SERVICIOS J&G', 'BANCO DE VENEZUELA')]:
            p = Park(keys={(agent, b, '001234561', '1')}, afipos={'001234561'})
            out = self.classify([record(agent, b), record(agent, b, afipos='999'), record(agent, b, NUMPOS='2')], p)
            self.assertEqual(out.FACTURACION.tolist(), [MODES[0], MODES[1], MODES[3]])
        p = Park(keys={('SERVICIOS J&G', 'BANCO DE VENEZUELA', '001234561', '1')}, afipos={'001234561'})
        self.assertEqual(self.classify([record('RAPIDPAGO')], p).FACTURACION.iloc[0], MODES[3])

    def test_missing_park_never_silently_proforma(self):
        out = self.classify([record('RAPIDPAGO')], Park(usable=False, reason='Falta parque'))
        self.assertEqual(out.FACTURACION.iloc[0], MODES[3])

    def test_exclusions_exact_not_substring(self):
        rows = [record(a) for a in ['CARROPAGO', 'FLEXIPOS', 'ACIERTO TECNOLOGICO', 'AGILPAGOS', 'CARROPAGO NUEVO']]
        out, excluded = classify(pd.DataFrame(rows), MAPPING, set(), Park())
        self.assertEqual(len(out), 1)
        self.assertEqual(sum(excluded.values()), 4)
        self.assertEqual(out.PERTENENCIA.iloc[0], 'CARROPAGO NUEVO')

    def test_unknown_and_bad_bank_pending(self):
        out = self.classify([record('SIN DICCIONARIO'), record(CODIGO_BANCO='896.0'), record(afipos='')])
        self.assertEqual(out.FACTURACION.tolist(), [MODES[3]]*3)

    def test_bank_variants_and_identifiers(self):
        self.assertEqual(bank('896.0'), 'BANFANB')
        self.assertEqual(bank('banfab'), 'BANFANB')
        self.assertEqual(identifier('00123.0'), '00123')
        self.assertEqual(identifier('29600100123020303913'), '29600100123020303913')
        with self.assertRaises(ValueError):
            identifier(29600100123020303913)


class InputTests(unittest.TestCase):
    def test_dictionary_normalization_and_errors(self):
        m, invalid = load_dictionary(excel([[None, 'RNC34', 'CORRECTO'], [None, '  rápido  ', 'RAPIDPAGO'], [None, 'OFICINA', '#N/A']]))
        self.assertEqual(m['RAPIDO'], 'RAPIDPAGO')
        self.assertEqual(m['RAPIDPAGO'], 'RAPIDPAGO')
        self.assertEqual(invalid, ['OFICINA'])
        with self.assertRaises(ValueError):
            load_dictionary(excel([['RNC34', 'CORRECTO'], ['A', 'B'], ['A', 'C']]))

    def test_hierarchy_and_duplicates(self):
        pairs, duplicates = load_billing(excel([['VEPAGOS'], ['MI BANCO'], ['MI BANCO'], ['RAPIDPAGO'], ['BANFANB']]), MAPPING)
        self.assertEqual(duplicates, 1)
        self.assertEqual(pairs, {('VEPAGOS', 'MI BANCO'), ('RAPIDPAGO', 'BANFANB')})
        with self.assertRaises(ValueError):
            load_billing(excel([['MI BANCO']]), MAPPING)
        with self.assertRaises(ValueError):
            load_billing(excel([['AGENTE DESCONOCIDO'], ['MI BANCO']]), MAPPING)

    def test_park_codes_and_deduplication(self):
        headers = ['BANCO', 'AGREGADO', 'AFIPOS', 'TERMINAL']
        row = [896, 'RAPID PAGO', '001234561', 1]
        p = load_park(excel([headers, row, row], '7A_PARQUE_PARCIAL_ACTIVOS'), MAPPING)
        self.assertEqual(len(p.keys), 1)
        self.assertEqual(p.duplicates, 1)
        self.assertIn(('RAPIDPAGO', 'BANFANB', '001234561', '1'), p.keys)

    def test_park_period(self):
        self.assertTrue(park_period_issue('Parque parcial Septiembre.xlsx', '2026-10'))
        self.assertFalse(park_period_issue('Parque octubre.xlsx', '2026-10'))
        self.assertTrue(park_period_issue('Parque 2025-10.xlsx', '2026-10'))


class IntegrationTests(unittest.TestCase):
    def inputs(self):
        d = excel([['RNC34', 'CORRECTO']] + list(MAPPING.items()))
        b = excel([['VEPAGOS'], ['BANCO DE VENEZUELA']])
        p = excel([['BANCO', 'AGREGADO', 'AFIPOS', 'TERMINAL']], '7A_PARQUE_PARCIAL_ACTIVOS')
        return d, b, p

    def run_process(self, rows, output, period='2026-10'):
        csv = io.BytesIO(pd.DataFrame(rows).to_csv(index=False, sep=';').encode('utf-8'))
        return process(csv, *self.inputs(), period, output, chunksize=2)

    def test_export_reconciles_and_preserves_text_duplicates_small_agent(self):
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp)/'out.xlsx'
            rows = [record(), record(), record('CARROPAGO'), record('PEQUENO'), record('SIN DICCIONARIO')]
            result = self.run_process(rows, out)
            self.assertEqual((result.audit.read, result.audit.processed), (5, 4))
            self.assertEqual(result.audit.duplicate_keys, 4)
            self.assertEqual(int(result.summary.EQUIPOS.sum()), 4)
            w = openpyxl.load_workbook(out, data_only=True)
            self.assertEqual(w.sheetnames, ['RNC', 'RESUMEN'])
            self.assertEqual(w['RNC'].max_row, 5)
            self.assertEqual(w['RNC']['A2'].value, '001234561')
            self.assertEqual(w['RNC']['F2'].value, '29600100123020303913')
            totals = [r[1].value for r in w['RESUMEN'].iter_rows() if r[0].value == 'Total general']
            self.assertEqual(totals, [4])
            self.assertEqual(len(w['RESUMEN'].data_validations.dataValidation), 4)
            self.assertTrue(w['RNC'].auto_filter.ref)
            self.assertTrue(w['RESUMEN'].auto_filter.ref)
            self.assertIn('PEQUENO', [r[0].value for r in w['RESUMEN'].iter_rows()])
            w.close()
            # A second upload replaces all counts, not an accumulated prior result.
            result = self.run_process([record('PEQUENO', afipos='999')], out)
            self.assertEqual(int(result.summary.EQUIPOS.sum()), 1)
            self.assertEqual(set(result.summary.AGENTE), {'PEQUENO'})

    def test_period_mismatch_stops_export(self):
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp)/'out.xlsx'
            with self.assertRaisesRegex(ValueError, 'otro período'):
                self.run_process([record()], out, '2026-09')
            self.assertFalse(out.exists())

    def test_all_excluded_produces_zero_totals(self):
        with tempfile.TemporaryDirectory() as temp:
            result = self.run_process([record('CARROPAGO')], Path(temp)/'out.xlsx')
            self.assertEqual(result.audit.processed, 0)
            self.assertEqual(len(result.summary), 0)

    def test_formula_injection_is_text(self):
        with tempfile.TemporaryDirectory() as temp:
            out = Path(temp)/'out.xlsx'
            self.run_process([record(SERIAL='=1+1')], out)
            w = openpyxl.load_workbook(out)
            self.assertEqual(w['RNC']['F2'].data_type, 's')
            w.close()


if __name__ == '__main__':
    unittest.main()
