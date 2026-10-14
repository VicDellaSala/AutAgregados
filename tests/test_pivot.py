import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import openpyxl
import test_processing as fixtures
from agregados.normalizacion import COLUMNS

N={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

class NativePivotTests(unittest.TestCase):
    def test_native_cache_contains_only_current_equipment(self):
        with tempfile.TemporaryDirectory() as temp:
            output=Path(temp)/'out.xlsx'
            fixtures.IntegrationTests().run_process([fixtures.record(), fixtures.record(), fixtures.record('PEQUENO'), fixtures.record('CARROPAGO')], output)
            with ZipFile(output) as z:
                p=ET.fromstring(z.read('xl/pivotTables/pivotTable1.xml'))
                c=ET.fromstring(z.read('xl/pivotCache/pivotCacheDefinition1.xml'))
                records=ET.fromstring(z.read('xl/pivotCache/pivotCacheRecords1.xml'))
                self.assertEqual(c.find('s:cacheSource/s:worksheetSource',N).attrib, {'ref':'A1:M4','sheet':'RNC'})
                self.assertEqual(int(c.get('recordCount')),3)
                self.assertEqual(len(records),3)
                self.assertTrue(all(len(r)==13 for r in records))
                self.assertEqual([f.get('name') for f in c.find('s:cacheFields',N)], COLUMNS)
                self.assertEqual(p.get('name'),'TablaDinámica1')
                self.assertEqual(p.find('s:dataFields/s:dataField',N).get('subtotal'),'count')
                self.assertEqual([f.get('x') for f in p.find('s:rowFields',N)],['8'])
                self.assertEqual([f.get('fld') for f in p.find('s:pageFields',N)],['11'])
                for name in ['xl/worksheets/sheet1.xml','xl/worksheets/sheet2.xml']:
                    root=ET.fromstring(z.read(name))
                    self.assertIsNone(root.find('.//s:pane',N))
                    self.assertIsNone(root.find('s:sheetProtection',N))
            w=openpyxl.load_workbook(output)
            self.assertEqual(w.sheetnames,['Hoja1','RNC'])
            self.assertEqual(len(w['Hoja1']._pivots),1)
            self.assertEqual(w['Hoja1']['B2'].value,'PROFORMA')
            self.assertEqual(w['Hoja1'].max_column,2)
            self.assertLess(w['Hoja1'].max_row,10)
            self.assertEqual([c.value for c in w['RNC'][1]],COLUMNS)
            w.close()
            fixtures.IntegrationTests().run_process([fixtures.record('PEQUENO')],output)
            with ZipFile(output) as z:
                r=ET.fromstring(z.read('xl/pivotCache/pivotCacheRecords1.xml'))
                self.assertEqual(len(r),1)
                self.assertNotIn(b'VEPAGOS',z.read('xl/pivotCache/pivotCacheDefinition1.xml'))

    def test_no_invalid_empty_pivot(self):
        with tempfile.TemporaryDirectory() as temp:
            output=Path(temp)/'out.xlsx'
            fixtures.IntegrationTests().run_process([fixtures.record('CARROPAGO')],output)
            with ZipFile(output) as z:
                self.assertFalse(any(n.startswith('xl/pivotTables/') for n in z.namelist()))
